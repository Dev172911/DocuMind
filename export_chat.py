from docx import Document
from io import BytesIO

def export_single_as_docx(session):
    doc = Document()
    doc.add_heading("DocuMind - Chat Export", 0)
    doc.add_paragraph(f"Topic: {session['title']}")
    doc.add_paragraph("")

    for msg in session["history"]:
        if msg["role"] == "user":
            doc.add_heading("You:", level=2)
        else:
            doc.add_heading("DocuMind:", level=2)
        doc.add_paragraph(msg["content"])
        doc.add_paragraph("")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def export_all_as_docx(sessions):
    doc = Document()
    doc.add_heading("DocuMind - Full Chat History", 0)

    for i, session in enumerate(sessions):
        doc.add_heading(f"Chat {i+1}: {session['title']}", level=1)
        for msg in session["history"]:
            if msg["role"] == "user":
                doc.add_heading("You:", level=2)
            else:
                doc.add_heading("DocuMind:", level=2)
            doc.add_paragraph(msg["content"])
            doc.add_paragraph("")
        doc.add_page_break()

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
